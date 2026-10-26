"""Deck validation: made-up numbers and overflowing text must never reach a slide."""
import copy

import pytest
from pptx import Presentation

from casebook.deck import DeckError, draft_deck, render_pptx, validate_deck

from conftest import FRAUD, LENDING, FakeLLM

GOOD = {
    "deck_title": "Brightline for Acme Bank",
    "deck_subtitle": "Fraud and lending platforms we have shipped",
    "cases": [
        {
            "corpus_title": FRAUD,
            "headline": "Streaming fraud scoring that cut chargebacks",
            "challenge": "A rules engine flagged too many good payments and missed new attacks.",
            "metrics": [
                {"value": "38%", "caption": "fewer chargebacks"},
                {"value": "42 ms", "caption": "p99 scoring latency"},
                {"value": "11,000", "caption": "weekly alerts, down from 40,000"},
            ],
            "delivered": ["Kafka streaming pipeline", "Gradient-boosted scoring model",
                          "Analyst review tool with weekly retraining"],
        },
        {
            "corpus_title": LENDING,
            "headline": "A custom loan platform that replaced a vendor",
            "challenge": "Expensive vendor licensing and manual re-keying of borrower data.",
            "metrics": [
                {"value": "45%", "caption": "lower annual platform cost"},
                {"value": "3 days", "caption": "to decision, down from 9"},
                {"value": "3", "caption": "systems with no manual re-entry"},
            ],
            "delivered": ["Temporal loan workflows", "Direct credit bureau integrations",
                          "Admin console for underwriting steps"],
        },
    ],
    "note": "",
}


def test_valid_deck_passes(corpus):
    assert validate_deck(GOOD, corpus) == []


def test_number_not_in_source_case_is_caught(corpus):
    deck = copy.deepcopy(GOOD)
    deck["cases"][0]["metrics"][0]["value"] = "45%"  # real number, but from the OTHER case
    problems = validate_deck(deck, corpus)
    assert len(problems) == 1 and "'45'" in problems[0] and FRAUD in problems[0]


def test_numbers_match_whole_not_as_substrings(corpus):
    deck = copy.deepcopy(GOOD)
    # "40,000" is in the source text, but "40" on its own is not.
    deck["cases"][0]["metrics"][0]["value"] = "40%"
    assert any("'40'" in p for p in validate_deck(deck, corpus))
    # "38%" is in the text, so a rounded "37%" must fail.
    deck["cases"][0]["metrics"][0]["value"] = "37%"
    assert any("'37'" in p for p in validate_deck(deck, corpus))


def test_invented_case_is_caught(corpus):
    deck = copy.deepcopy(GOOD)
    deck["cases"][0]["corpus_title"] = "Fraud Platform for Acme Bank"
    assert any("not a case study in the library" in p for p in validate_deck(deck, corpus))


def test_text_too_long_for_its_box_is_caught(corpus):
    deck = copy.deepcopy(GOOD)
    deck["cases"][1]["metrics"][2]["value"] = "Three systems"
    deck["deck_title"] = "x" * 60
    problems = validate_deck(deck, corpus)
    assert any("metric 3 value" in p for p in problems)
    assert any("deck_title is 60 chars" in p for p in problems)


def test_wrong_number_of_metrics_and_duplicate_cases_are_caught(corpus):
    deck = copy.deepcopy(GOOD)
    deck["cases"][1] = copy.deepcopy(deck["cases"][0])
    deck["cases"][1]["metrics"].pop()
    problems = validate_deck(deck, corpus)
    assert any("chosen twice" in p for p in problems)
    assert any("has 2 metrics" in p for p in problems)


def test_subtitle_number_must_be_supported(corpus):
    deck = copy.deepcopy(GOOD)
    deck["deck_subtitle"] = "How we cut fraud losses by 70%"
    assert any("deck_subtitle states '70'" in p for p in validate_deck(deck, corpus))


def test_failed_draft_gets_one_repair_round(corpus):
    bad = copy.deepcopy(GOOD)
    bad["cases"][0]["metrics"][0]["value"] = "61%"
    llm = FakeLLM(bad, GOOD)
    draft = draft_deck("Acme Bank wants to cut fraud", corpus, llm)
    assert draft.attempts == 2
    repair_request = llm.calls[1]["messages"][-1]["content"]
    assert "'61'" in repair_request


def test_deck_that_never_validates_raises_with_problems(corpus):
    bad = copy.deepcopy(GOOD)
    bad["cases"][0]["metrics"][0]["value"] = "61%"
    with pytest.raises(DeckError) as e:
        draft_deck("Acme Bank", corpus, FakeLLM(bad, bad))
    assert any("'61'" in p for p in e.value.problems)


def test_no_matching_cases_raises_with_the_models_reason(corpus):
    llm = FakeLLM({"cases": [], "note": "Nothing in the library covers satellite imaging."})
    with pytest.raises(DeckError, match="satellite"):
        draft_deck("satellite imaging startup", corpus, llm)


def test_render_produces_title_case_and_sources_slides(corpus, tmp_path):
    path = render_pptx(GOOD, corpus, str(tmp_path / "pitch.pptx"))
    prs = Presentation(path)
    assert len(prs.slides) == 4
    all_text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes
                        if sh.has_text_frame)
    assert "Brightline for Acme Bank" in all_text and "Sources" in all_text
